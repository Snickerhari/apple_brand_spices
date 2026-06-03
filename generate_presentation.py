import sys
import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # -------------------------------------------------------------
    # COLOR PALETTE (Premium Dark Theme)
    # -------------------------------------------------------------
    COLOR_BG = RGBColor(0x0F, 0x17, 0x2A)       # Very dark slate blue (#0F172A)
    COLOR_CARD_BG = RGBColor(0x1E, 0x29, 0x3B)  # Dark slate gray (#1E293B)
    COLOR_CARD_BG_ALT = RGBColor(0x33, 0x41, 0x55) # Slate gray highlight (#334155)
    COLOR_ACCENT = RGBColor(0xFB, 0xBF, 0x24)   # Warm Gold / Amber (#FBBF24)
    COLOR_TEXT_TITLE = RGBColor(0xF8, 0xFA, 0xFC) # White (#F8FAFC)
    COLOR_TEXT_BODY = RGBColor(0xCB, 0xD5, 0xE1)  # Light gray (#CBD5E1)
    COLOR_TEXT_MUTED = RGBColor(0x94, 0xA3, 0xB8) # Muted gray (#94A3B8)
    COLOR_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    
    FONT_TITLE = "Segoe UI"
    FONT_BODY = "Segoe UI"
    
    # -------------------------------------------------------------
    # HELPER FUNCTIONS
    # -------------------------------------------------------------
    def apply_bg(slide):
        """Sets a solid dark background for the slide."""
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_BG

    def add_slide_header(slide, title_text, category_text="SOCIOLOGY OF STRATIFICATION"):
        """Adds a unified header with category, title, and thin accent line."""
        # 1. Category Tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = FONT_BODY
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT
        
        # 2. Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(32)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT_TITLE
        
        # 3. Accent Line
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 
            Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.03)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_ACCENT
        line.line.fill.background() # borderless

    def add_card(slide, left, top, width, height, title_text="", body_paragraphs=None, is_highlight=False):
        """Creates a modern rounded card shape with title and paragraph content."""
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
        )
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG_ALT if is_highlight else COLOR_CARD_BG
        card.line.color.rgb = COLOR_ACCENT if is_highlight else COLOR_CARD_BG
        card.line.width = Pt(1.5) if is_highlight else Pt(1)
        
        # Inner text frame
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)
        
        # Title
        p_title = tf.paragraphs[0]
        if title_text:
            p_title.text = title_text
            p_title.font.name = FONT_TITLE
            p_title.font.size = Pt(18)
            p_title.font.bold = True
            p_title.font.color.rgb = COLOR_ACCENT if not is_highlight else COLOR_WHITE
            p_title.space_after = Pt(8)
            
        # Body
        if body_paragraphs:
            for idx, text in enumerate(body_paragraphs):
                if not title_text and idx == 0:
                    p = p_title
                else:
                    p = tf.add_paragraph()
                p.text = text
                p.font.name = FONT_BODY
                p.font.size = Pt(14)
                p.font.color.rgb = COLOR_TEXT_TITLE if is_highlight else COLOR_TEXT_BODY
                p.space_after = Pt(6)
        return card

    # Blank layout for custom positioning
    blank_slide_layout = prs.slide_layouts[6]

    # =============================================================
    # SLIDE 1: TITLE SLIDE (Premium & Minimalist)
    # =============================================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide1)
    
    # Elegant gold border outline on title slide
    border = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), Inches(0.4), Inches(12.533), Inches(6.7))
    border.fill.background()
    border.line.color.rgb = COLOR_ACCENT
    border.line.width = Pt(1.5)
    
    # Subtitle Category
    box_sub = slide1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(0.5))
    tf_sub = box_sub.text_frame
    p_sub = tf_sub.paragraphs[0]
    p_sub.text = "SOCIOLOGICAL STUDY & ANALYSIS  •  16 MARKS SPECIAL"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(12)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_ACCENT
    p_sub.alignment = PP_ALIGN.CENTER
    
    # Main Title
    box_main = slide1.shapes.add_textbox(Inches(1.5), Inches(2.8), Inches(10.333), Inches(2.0))
    tf_main = box_main.text_frame
    tf_main.word_wrap = True
    p_main = tf_main.paragraphs[0]
    p_main.text = "SOCIAL CLASS"
    p_main.font.name = FONT_TITLE
    p_main.font.size = Pt(56)
    p_main.font.bold = True
    p_main.font.color.rgb = COLOR_TEXT_TITLE
    p_main.alignment = PP_ALIGN.CENTER
    
    # Lower Title Line
    p_submain = tf_main.add_paragraph()
    p_submain.text = "Meaning, Definitions & Characteristics"
    p_submain.font.name = FONT_BODY
    p_submain.font.size = Pt(22)
    p_submain.font.color.rgb = COLOR_TEXT_BODY
    p_submain.alignment = PP_ALIGN.CENTER
    p_submain.space_before = Pt(10)
    
    # Bottom details
    box_details = slide1.shapes.add_textbox(Inches(1.5), Inches(5.2), Inches(10.333), Inches(1.0))
    tf_details = box_details.text_frame
    p_det = tf_details.paragraphs[0]
    p_det.text = "A Comprehensive Academic Resource on Social Stratification"
    p_det.font.name = FONT_BODY
    p_det.font.size = Pt(13)
    p_det.font.color.rgb = COLOR_TEXT_MUTED
    p_det.alignment = PP_ALIGN.CENTER

    # =============================================================
    # SLIDE 2: INTRODUCTION
    # =============================================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide2)
    add_slide_header(slide2, "Introduction to Social Class")
    
    # Left Hero Text
    box_left = slide2.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(5.2), Inches(4.5))
    tf_left = box_left.text_frame
    tf_left.word_wrap = True
    p_l1 = tf_left.paragraphs[0]
    p_l1.text = "Understanding Stratification"
    p_l1.font.name = FONT_TITLE
    p_l1.font.size = Pt(24)
    p_l1.font.bold = True
    p_l1.font.color.rgb = COLOR_ACCENT
    p_l1.space_after = Pt(14)
    
    p_l2 = tf_left.add_paragraph()
    p_l2.text = "Social class is one of the foundational concepts of modern sociology.\n\nIt describes how human societies organize themselves into distinct, horizontal layers. Unlike rigid pre-modern stratification systems, social class represents an open system directly tied to economic structures."
    p_l2.font.name = FONT_BODY
    p_l2.font.size = Pt(16)
    p_l2.font.color.rgb = COLOR_TEXT_BODY
    p_l2.space_after = Pt(14)
    
    # Right Side Card (Focal Point)
    right_card_text = [
        "• System of Division: Refers to the grouping of individuals based on shared socioeconomic factors.",
        "• Determinant Factors: Stratified based on social status, income, occupation, education, wealth, and lifestyle choices.",
        "• Opportunities & Positions: Members of the same social class generally share similar opportunities, resources, life chances, and social positions in society."
    ]
    add_card(
        slide2, 
        left=Inches(6.5), 
        top=Inches(2.0), 
        width=Inches(6.0), 
        height=Inches(4.5), 
        title_text="Key Dimensions of Class", 
        body_paragraphs=right_card_text, 
        is_highlight=True
    )

    # =============================================================
    # SLIDE 3: SOCIOLOGICAL DEFINITIONS
    # =============================================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide3)
    add_slide_header(slide3, "Sociological Definitions")
    
    # Three Column Cards for major sociological thinkers
    card_width = Inches(3.64)
    card_height = Inches(4.5)
    gap = Inches(0.4)
    start_left = Inches(0.8)
    top_pos = Inches(2.0)
    
    # Definition 1: MacIver and Page
    add_card(
        slide3, 
        left=start_left, 
        top=top_pos, 
        width=card_width, 
        height=card_height, 
        title_text="MacIver and Page",
        body_paragraphs=[
            "“A social class is any portion of a community marked off from the rest by social status.”",
            "",
            "Key Focus: Status stratification. Society distinguishes groups based on prestige, social honor, and relative position."
        ]
    )
    
    # Definition 2: Ogburn and Nimkoff
    add_card(
        slide3, 
        left=start_left + card_width + gap, 
        top=top_pos, 
        width=card_width, 
        height=card_height, 
        title_text="Ogburn & Nimkoff",
        body_paragraphs=[
            "“A social class is the aggregate of persons having essentially the same social status in a given society.”",
            "",
            "Key Focus: Collective status. The focus is on the aggregate grouping of individuals sharing a common status pool."
        ]
    )
    
    # Definition 3: Max Weber (Highlight/Accent Card)
    add_card(
        slide3, 
        left=start_left + (card_width + gap) * 2, 
        top=top_pos, 
        width=card_width, 
        height=card_height, 
        title_text="Max Weber",
        body_paragraphs=[
            "“A social class consists of people who share similar economic opportunities and life chances.”",
            "",
            "Key Focus: Market power & life chances. Weber emphasizes the economic foundation of class positions."
        ],
        is_highlight=True
    )

    # =============================================================
    # SLIDE 4: MEANING & SYSTEM STRUCTURE
    # =============================================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide4)
    add_slide_header(slide4, "Meaning of Social Class")
    
    # Left Box - Grouping description
    add_card(
        slide4, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(5.6), 
        height=Inches(4.5), 
        title_text="Group Status & Position",
        body_paragraphs=[
            "• Social Position: A social class is a group of people who occupy a similar position in the social hierarchy.",
            "",
            "• Strata Identity: Members of a class share an objective position that shapes their access to material resources, culture, and influence.",
            "",
            "• Contrast to Caste: Unlike caste, which is a closed system based strictly on birth, social class is fluid, achievement-based, and economically defined."
        ]
    )
    
    # Right Box - Mobility & Achievement
    add_card(
        slide4, 
        left=Inches(6.9), 
        top=Inches(2.0), 
        width=Inches(5.6), 
        height=Inches(4.5), 
        title_text="Achievement-Based Stratification",
        body_paragraphs=[
            "• Earned Stratum: Status is not biologically or legally fixed. It is heavily influenced by personal efforts, merits, and choices.",
            "",
            "• Social Mobility: Individuals can move vertically (upwards or downwards) between different classes throughout their lifetime.",
            "",
            "• Pathways of Movement: Mobility is primarily achieved through: \n  - Education and skills acquisition\n  - Occupational advancement\n  - Wealth accumulation & hard work"
        ],
        is_highlight=True
    )

    # =============================================================
    # SLIDE 5: CHARACTERISTICS (1 - 4)
    # =============================================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide5)
    add_slide_header(slide5, "Characteristics of Social Class (1 - 4)")
    
    # 2x2 Grid Layout
    w_card = Inches(5.6)
    h_card = Inches(2.1)
    g_w = Inches(0.533)
    g_h = Inches(0.3)
    l_start = Inches(0.8)
    t_start = Inches(2.0)
    
    # Characteristic 1
    add_card(
        slide5, 
        left=l_start, 
        top=t_start, 
        width=w_card, 
        height=h_card, 
        title_text="1. Hierarchical Division",
        body_paragraphs=["Society is divided into distinct, vertical classes such as the upper, middle, and lower classes. Each class occupies a specific rank or tier relative to others."]
    )
    
    # Characteristic 2
    add_card(
        slide5, 
        left=l_start + w_card + g_w, 
        top=t_start, 
        width=w_card, 
        height=h_card, 
        title_text="2. Social Status & Prestige",
        body_paragraphs=["Each class carries a distinct level of prestige, status, and social honor. Behavior, speech, and consumption patterns reinforce this status."]
    )
    
    # Characteristic 3
    add_card(
        slide5, 
        left=l_start, 
        top=t_start + h_card + g_h, 
        width=w_card, 
        height=h_card, 
        title_text="3. Open Group System",
        body_paragraphs=["Class operates as an open stratification system. Individuals are not bound to the class of their birth; they are free to transition to other strata based on achievements."]
    )
    
    # Characteristic 4
    add_card(
        slide5, 
        left=l_start + w_card + g_w, 
        top=t_start + h_card + g_h, 
        width=w_card, 
        height=h_card, 
        title_text="4. Class Consciousness",
        body_paragraphs=["Members of a class share awareness of their position, common interests, and collective status in relation to other social groups."]
    )

    # =============================================================
    # SLIDE 6: CHARACTERISTICS (5 - 8)
    # =============================================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide6)
    add_slide_header(slide6, "Characteristics of Social Class (5 - 8)")
    
    # Characteristic 5
    add_card(
        slide6, 
        left=l_start, 
        top=t_start, 
        width=w_card, 
        height=h_card, 
        title_text="5. Similar Lifestyle",
        body_paragraphs=["People within a class share similar lifestyles, consumer habits, tastes, standard of living, and recreational patterns."]
    )
    
    # Characteristic 6
    add_card(
        slide6, 
        left=l_start + w_card + g_w, 
        top=t_start, 
        width=w_card, 
        height=h_card, 
        title_text="6. Based on Achievement",
        body_paragraphs=["Class is primarily determined by achieved status variables like education, occupation, income, and accumulated wealth rather than ascribed birth."]
    )
    
    # Characteristic 7
    add_card(
        slide6, 
        left=l_start, 
        top=t_start + h_card + g_h, 
        width=w_card, 
        height=h_card, 
        title_text="7. Social Mobility",
        body_paragraphs=["Vertical and horizontal movement is structurally supported. Upward mobility is fueled by education, employment, and economic success."]
    )
    
    # Characteristic 8
    add_card(
        slide6, 
        left=l_start + w_card + g_w, 
        top=t_start + h_card + g_h, 
        width=w_card, 
        height=h_card, 
        title_text="8. No Legal/Religious Restrictions",
        body_paragraphs=["Unlike caste or estate systems, class has no legal or religious barriers. There are no restrictions on inter-group interactions or marriages."]
    )

    # =============================================================
    # SLIDE 7: TYPES OF SOCIAL CLASSES
    # =============================================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide7)
    add_slide_header(slide7, "Three-Tier Structure of Social Class")
    
    # 3 Column structure (Upper, Middle, Lower)
    c_w = Inches(3.64)
    c_h = Inches(4.5)
    c_g = Inches(0.4)
    c_l = Inches(0.8)
    c_t = Inches(2.0)
    
    # Upper Class
    add_card(
        slide7, 
        left=c_l, 
        top=c_t, 
        width=c_w, 
        height=c_h, 
        title_text="Upper Class",
        body_paragraphs=[
            "• Composition: Consists of wealthy industrialists, large business owners, and land elites.",
            "",
            "• Economic Basis: High capital accumulation, asset ownership, and generational wealth.",
            "",
            "• Social Impact: Wields major economic and political influence, enjoying maximum life opportunities and elite status."
        ]
    )
    
    # Middle Class (Highlight)
    add_card(
        slide7, 
        left=c_l + c_w + c_g, 
        top=c_t, 
        width=c_w, 
        height=c_h, 
        title_text="Middle Class",
        body_paragraphs=[
            "• Composition: Comprises professionals, teachers, engineers, corporate and government employees.",
            "",
            "• Economic Basis: Salary-based income, professional education, and skilled expertise.",
            "",
            "• Social Impact: Act as the stabilizing core of modern society. Characterized by high valuation of education and social mobility."
        ],
        is_highlight=True
    )
    
    # Lower Class
    add_card(
        slide7, 
        left=c_l + (c_w + c_g) * 2, 
        top=c_t, 
        width=c_w, 
        height=c_h, 
        title_text="Lower Class",
        body_paragraphs=[
            "• Composition: Includes manual laborers, daily wage earners, agricultural workers, and economically weaker sections.",
            "",
            "• Economic Basis: Low-wage labor, lack of stable assets, and economic insecurity.",
            "",
            "• Social Impact: Struggles with limited life chances, lower educational access, and higher exposure to social vulnerability."
        ]
    )

    # =============================================================
    # SLIDE 8: IMPORTANCE OF SOCIAL CLASS
    # =============================================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide8)
    add_slide_header(slide8, "Sociological Importance of Social Class")
    
    # Left Hero card
    add_card(
        slide8, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(5.0), 
        height=Inches(4.5), 
        title_text="Why Class Analysis Matters",
        body_paragraphs=[
            "Social class is not just a statistical grouping. It represents an active force that influences every aspect of human life.",
            "",
            "It affects child development, educational outcomes, physical health, psychological well-being, lifestyle, and final life expectancy.",
            "",
            "Analyzing social class reveals how societal rewards, power, and privilege are distributed."
        ],
        is_highlight=True
    )
    
    # Right Side List
    box_r = slide8.shapes.add_textbox(Inches(6.3), Inches(2.0), Inches(6.2), Inches(4.5))
    tf_r = box_r.text_frame
    tf_r.word_wrap = True
    
    pts = [
        ("Understanding Social Inequality", "It serves as a primary tool to measure and map economic inequalities, resource distribution, and structural imbalances within a society."),
        ("Predicting Life Chances", "Class position heavily determines an individual's access to quality education, healthcare, legal protection, and general security."),
        ("Shaping Lifestyles & Interaction", "Class determines consumption habits, leisure activities, residential patterns, and limits or facilitates social network connections."),
        ("Political Dynamics & Voting Behavior", "Power structures are deeply linked to class. It dictates political participation, interest lobbies, and patterns of ideological voting.")
    ]
    
    for idx, (title, desc) in enumerate(pts):
        p_t = tf_r.paragraphs[0] if idx == 0 else tf_r.add_paragraph()
        p_t.text = f"{idx+1}. {title}"
        p_t.font.name = FONT_TITLE
        p_t.font.size = Pt(16)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_ACCENT
        p_t.space_before = Pt(8) if idx > 0 else Pt(0)
        
        p_d = tf_r.add_paragraph()
        p_d.text = desc
        p_d.font.name = FONT_BODY
        p_d.font.size = Pt(13)
        p_d.font.color.rgb = COLOR_TEXT_BODY
        p_d.space_after = Pt(10)

    # =============================================================
    # SLIDE 9: CLASS SYSTEM VS CASTE SYSTEM (COMPREHENSIVE TABLE)
    # =============================================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide9)
    add_slide_header(slide9, "Class System vs. Caste System")
    
    # Create Table shape
    rows = 9
    cols = 3
    left_tbl = Inches(0.8)
    top_tbl = Inches(2.0)
    width_tbl = Inches(11.733)
    height_tbl = Inches(4.5)
    
    table_shape = slide9.shapes.add_table(rows, cols, left_tbl, top_tbl, width_tbl, height_tbl)
    table = table_shape.table
    
    # Set column widths
    table.columns[0].width = Inches(2.2)  # Feature
    table.columns[1].width = Inches(4.766) # Class
    table.columns[2].width = Inches(4.766) # Caste
    
    # Content arrays
    headers = ["Basis / Feature", "Class System", "Caste System"]
    data = [
        ["1. Foundation", "Based on wealth, education, occupation, and income.", "Based strictly on birth, ancestry, and heredity."],
        ["2. System Type", "Open system of social stratification.", "Closed, rigid system of stratification."],
        ["3. Social Mobility", "Social mobility is common and structured.", "Social mobility is strictly prohibited/non-existent."],
        ["4. Status Type", "Status is achieved through personal efforts.", "Status is ascribed by birth and unchangeable."],
        ["5. Marriage", "No restrictions; inter-class marriage is normal.", "Endogamous; marriage restricted within the caste."],
        ["6. Occupation", "Free choice; occupation is fluid and non-fixed.", "Traditionally fixed; hereditary occupational roles."],
        ["7. Social Context", "Predominantly found in modern industrial societies.", "Traditionally rooted in historic Indian society."],
        ["8. Relations", "Flexible, open, and dynamic social interactions.", "Restricted, rigid social relations & pollution norms."]
    ]
    
    # Apply styling & fill headers
    for col_idx, text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG_ALT
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.name = FONT_TITLE
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.alignment = PP_ALIGN.CENTER
        
    # Populate & style data rows
    for row_idx, row_data in enumerate(data):
        row_num = row_idx + 1
        # Alternate background colors for table readability
        bg_color = COLOR_CARD_BG if row_num % 2 == 1 else COLOR_BG
        
        for col_idx, cell_value in enumerate(row_data):
            cell = table.cell(row_num, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_color
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            
            p = cell.text_frame.paragraphs[0]
            p.text = cell_value
            p.font.name = FONT_BODY
            p.font.size = Pt(12)
            
            # Left align first column, others left align but center headers
            if col_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_TEXT_TITLE
                p.alignment = PP_ALIGN.LEFT
            else:
                p.font.color.rgb = COLOR_TEXT_BODY
                p.alignment = PP_ALIGN.LEFT

    # =============================================================
    # SLIDE 10: CONCLUSION
    # =============================================================
    slide10 = prs.slides.add_slide(blank_slide_layout)
    apply_bg(slide10)
    add_slide_header(slide10, "Summary & Conclusion")
    
    # Left Hero Panel
    add_card(
        slide10, 
        left=Inches(0.8), 
        top=Inches(2.0), 
        width=Inches(5.0), 
        height=Inches(4.5), 
        title_text="Key Summary",
        body_paragraphs=[
            "• Social Class represents an open, achievement-based hierarchy determined by economic resources, occupation, and education.",
            "",
            "• Characteristics such as social mobility and the absence of legal barriers make it the defining system of stratification in modern democratic states.",
            "",
            "• Unlike the static caste system, social class encourages personal effort and systemic growth."
        ],
        is_highlight=True
    )
    
    # Right Takeaway Box
    add_card(
        slide10, 
        left=Inches(6.3), 
        top=Inches(2.0), 
        width=Inches(6.2), 
        height=Inches(4.5), 
        title_text="Academic Synthesis (16 Marks Focal Points)",
        body_paragraphs=[
            "1. Theoretical Alignment:",
            "Integrate Marx's production-based class model (bourgeoisie vs. proletariat) and Weber's multidimensional model (class, status, power) for top-tier exam answers.",
            "",
            "2. Open vs. Closed Systems:",
            "Emphasize the transition from traditional, closed systems (caste, estate) to modern, open class stratification systems driven by global industrialization.",
            "",
            "3. Societal Significance:",
            "Class position remains the single strongest predictor of individual life chances, educational attainment, health access, and political participation in contemporary societies."
        ]
    )

    # Save presentation
    filename = "Social_Class_Presentation.pptx"
    prs.save(filename)
    print(f"Presentation saved successfully as '{filename}'")

if __name__ == "__main__":
    create_presentation()
